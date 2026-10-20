#!/usr/bin/env python3
"""
blackslides - make PowerPoint slides (and figure images) ready for a black background.

What it does to each selected slide
  * sets the slide background to black (or --bg colour)
  * swaps the slide's theme colour mapping so inherited dark text becomes light
  * re-colours explicit text, line and fill colours: neutrals invert (black <-> white,
    dark grey <-> light grey), saturated colours keep their hue and stay bright
  * converts every embedded picture (PNG/JPEG/GIF/BMP/TIFF) the same way, pixel by pixel
  * SVG pictures are re-coloured as vectors (their PNG fallback is converted too)
  * EMF/WMF pictures are re-coloured record by record, so they stay vectors
  * nothing is resampled: pictures keep their exact pixel size, and converted
    pictures are saved losslessly (JPEGs become PNG unless --keep-jpeg)
  * leaves micrographs / photos with dark backgrounds alone (auto-detected)

You can also give it image files directly (png/jpg/tif/svg) instead of a .pptx.

Per-picture overrides: put "#keep" in a picture's Alt Text to leave it untouched, or
"#force" to convert it even if it looks like a dark photo.

Usage
  blackslides talk.pptx                    -> talk_black.pptx
  blackslides talk.pptx --slides 3,5-8     -> only those slides
  blackslides talk.pptx --brighten 0.3 --saturation 1.2
  blackslides fig1.png fig2.jpg            -> fig1_black.png, fig2_black.png
"""

import argparse
import base64
import colorsys
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass

import numpy as np
from PIL import Image, ImageColor
from lxml import etree

try:
    from scipy import ndimage
except ImportError:  # photo protection is skipped without scipy
    ndimage = None

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
RT_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
RT_THEME = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
RT_CHART = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart"
RT_DGM_DRAWING = "http://schemas.microsoft.com/office/2007/relationships/diagramDrawing"
R_EMBED = "{%s}embed" % R_NS


def a(tag):
    return "{%s}%s" % (A_NS, tag)


def local(el):
    return etree.QName(el).localname if isinstance(el.tag, str) else ""


# ----------------------------------------------------------------------------
# Colour science
# ----------------------------------------------------------------------------

@dataclass
class Settings:
    method: str = "oklab"      # oklab | hsl | invert
    protect: float = 1.0       # 0..1: how strongly saturated colours keep (bright) lightness
    brighten: float = 0.0      # 0..1: lift mid-tones after inversion
    saturation: float = 1.0    # chroma multiplier after inversion
    black_point: float = 0.0   # output lightness below this becomes pure black
    bg: tuple = (0, 0, 0)
    dpi: int = 300
    protect_photos: bool = True
    keep_jpeg: bool = False            # re-save JPEGs as JPEG (smaller, slightly lossy)
    rasterize_metafiles: bool = False  # EMF/WMF -> PNG instead of recolouring the vectors


_M1 = np.array([[0.4122214708, 0.5363325363, 0.0514459929],
                [0.2119034982, 0.6806995451, 0.1073969566],
                [0.0883024619, 0.2817188376, 0.6299787005]], dtype=np.float32)
_M2 = np.array([[0.2104542553, 0.7936177850, -0.0040720468],
                [1.9779984951, -2.4285922050, 0.4505937099],
                [0.0259040371, 0.7827717662, -0.8086757660]], dtype=np.float32)
_M2_INV = np.linalg.inv(_M2.astype(np.float64)).astype(np.float32)
_M1_INV = np.linalg.inv(_M1.astype(np.float64)).astype(np.float32)


def _to_linear(c):
    return np.where(c <= 0.04045, c / 12.92, ((c + 0.055) / 1.055) ** 2.4)


def _to_srgb(c):
    c = np.clip(c, 0.0, 1.0)
    return np.where(c <= 0.0031308, 12.92 * c, 1.055 * np.power(c, 1 / 2.4) - 0.055)


def rgb_to_oklab(rgb):
    return np.cbrt(_to_linear(rgb) @ _M1.T) @ _M2.T


def oklab_to_linear(lab):
    return ((lab @ _M2_INV.T) ** 3) @ _M1_INV.T


def _oklab_to_rgb_gamut(L, a_, b_):
    """OKLab -> sRGB, reducing chroma (keeping hue and lightness) when out of gamut."""
    lab = np.stack([L, a_, b_], axis=-1)
    lin = oklab_to_linear(lab)
    bad = np.any((lin < -1e-4) | (lin > 1 + 1e-4), axis=-1)
    if bad.any():
        Lb, ab, bb = L[bad], a_[bad], b_[bad]
        lo = np.zeros_like(Lb)
        hi = np.ones_like(Lb)
        for _ in range(14):
            mid = (lo + hi) / 2
            t = oklab_to_linear(np.stack([Lb, ab * mid, bb * mid], axis=-1))
            ok = np.all((t >= -1e-4) & (t <= 1 + 1e-4), axis=-1)
            lo = np.where(ok, mid, lo)
            hi = np.where(ok, hi, mid)
        lin[bad] = oklab_to_linear(np.stack([Lb, ab * lo, bb * lo], axis=-1))
    return _to_srgb(lin)


def transform_rgb(rgb, s):
    """rgb: float32 array (..., 3) in 0..1. Returns the black-background version."""
    rgb = rgb.astype(np.float32)
    if s.method == "invert":
        out = 1.0 - rgb
    elif s.method == "hsl":
        # HSL lightness inversion with hue & saturation kept (= invert + rotate hue 180)
        out = rgb + 1.0 - rgb.max(-1, keepdims=True) - rgb.min(-1, keepdims=True)
    else:
        out = None
    if out is not None and s.brighten == 0 and s.saturation == 1 and s.black_point == 0:
        return np.clip(out, 0, 1)

    lab = rgb_to_oklab(rgb if out is None else np.clip(out, 0, 1))
    L, a_, b_ = lab[..., 0], lab[..., 1], lab[..., 2]
    if out is None:
        chroma = np.hypot(a_, b_)
        w = np.clip(chroma / 0.12, 0, 1) * s.protect   # 0 = neutral grey, 1 = vivid colour
        # lightness of the RGB-inverted grey: grey 204 -> 51 exactly as a plain negative would
        L_inv = np.cbrt(_to_linear(1.0 - _to_srgb(np.clip(L, 0, 1) ** 3)))
        # neutrals invert; vivid colours take whichever of L / 1-L is brighter
        L = (1 - w) * L_inv + w * np.maximum(L, L_inv)
    if s.black_point > 0:
        L = np.clip((L - s.black_point) / (1 - s.black_point), 0, 1)
    if s.brighten > 0:
        L = np.power(np.clip(L, 0, 1), 1.0 - 0.6 * s.brighten)
    a_ = a_ * s.saturation
    b_ = b_ * s.saturation
    return _oklab_to_rgb_gamut(np.clip(L, 0, 1), a_, b_)


def transform_hex(hexval, s):
    rgb = np.array([[int(hexval[i:i + 2], 16) / 255 for i in (0, 2, 4)]], dtype=np.float32)
    out = np.round(transform_rgb(rgb, s)[0] * 255).astype(int)
    return "%02X%02X%02X" % tuple(out)


MIN_TEXT_L = 0.68  # OKLab lightness floor for text on a dark background


def lift_lightness(hexval, min_l):
    """Raise a colour's OKLab lightness to at least min_l, keeping its hue."""
    rgb = np.array([[int(hexval[i:i + 2], 16) / 255 for i in (0, 2, 4)]], dtype=np.float32)
    lab = rgb_to_oklab(rgb)
    if lab[0, 0] >= min_l:
        return hexval
    out = _oklab_to_rgb_gamut(np.array([min_l], np.float32), lab[:, 1], lab[:, 2])[0]
    return "%02X%02X%02X" % tuple(int(round(v * 255)) for v in out)


def rel_luminance(hexval):
    c = _to_linear(np.array([int(hexval[i:i + 2], 16) / 255 for i in (0, 2, 4)]))
    return float(0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2])


def contrast(h1, h2):
    l1, l2 = sorted((rel_luminance(h1), rel_luminance(h2)), reverse=True)
    return (l1 + 0.05) / (l2 + 0.05)


# ----------------------------------------------------------------------------
# Raster images
# ----------------------------------------------------------------------------

def _luma(rgb):
    return 0.299 * rgb[..., 0] + 0.587 * rgb[..., 1] + 0.114 * rgb[..., 2]


def find_dark_panels(rgb):
    """Boolean mask of rectangular dark-background panels (micrographs, photos)
    that should NOT be inverted. Solid black shapes are not counted."""
    h, w = rgb.shape[:2]
    keep = np.zeros((h, w), bool)
    if ndimage is None or min(h, w) < 40:
        return keep
    lum = _luma(rgb)
    k = max(9, int(0.03 * min(h, w)))
    density = ndimage.uniform_filter((lum < 0.3).astype(np.float32), size=k)
    dense = density > 0.8
    labels, n = ndimage.label(dense)
    for i, sl in enumerate(ndimage.find_objects(labels), start=1):
        ys, xs = sl
        bh, bw = ys.stop - ys.start, xs.stop - xs.start
        if bh * bw < 0.004 * h * w or min(bh, bw) < 2 * k:
            continue
        if (labels[sl] == i).mean() < 0.6:
            continue
        # grow the box back by the filter radius
        y0, y1 = max(0, ys.start - k // 2), min(h, ys.stop + k // 2)
        x0, x1 = max(0, xs.start - k // 2), min(w, xs.stop + k // 2)
        region = rgb[y0:y1, x0:x1]
        rl = lum[y0:y1, x0:x1]
        flat = (region * 255).astype(np.uint8).reshape(-1, 3)
        n_colours = len(np.unique(flat[:: max(1, len(flat) // 200000)], axis=0))
        # real images have texture/noise; solid black bars do not
        if rl.std() > 0.02 and n_colours > 64:
            keep[y0:y1, x0:x1] = True
    return keep


def is_dark_image(rgb, alpha=None):
    """True if the picture already has a dark background (look at its border)."""
    lum = _luma(rgb)
    b = max(1, int(0.02 * min(lum.shape)))
    border = np.concatenate([lum[:b].ravel(), lum[-b:].ravel(), lum[:, :b].ravel(), lum[:, -b:].ravel()])
    if alpha is not None:
        ab = np.concatenate([alpha[:b].ravel(), alpha[-b:].ravel(), alpha[:, :b].ravel(), alpha[:, -b:].ravel()])
        border = border[ab > 0.5]
        if border.size == 0:  # transparent border: decide from opaque content
            opaque = lum[alpha > 0.5]
            return opaque.size > 0 and np.median(opaque) < 0.25 and opaque.mean() < 0.35
    return np.median(border) < 0.2 and lum.mean() < 0.35


def transform_pil(img, s, force=False):
    """Returns (new PIL image, status) where status is converted / dark-skipped."""
    has_alpha = img.mode in ("RGBA", "LA", "PA") or (img.mode == "P" and "transparency" in img.info)
    img = img.convert("RGBA" if has_alpha else "RGB")
    arr = np.asarray(img)
    rgb = arr[..., :3].astype(np.float32) / 255
    alpha = arr[..., 3].astype(np.float32) / 255 if has_alpha else None
    if not force and s.protect_photos and is_dark_image(rgb, alpha):
        return None, "dark-skipped"
    keep = find_dark_panels(rgb) if (s.protect_photos and not force) else None

    # transform each distinct colour once
    u8 = arr[..., :3].reshape(-1, 3)
    packed = (u8[:, 0].astype(np.uint32) << 16) | (u8[:, 1].astype(np.uint32) << 8) | u8[:, 2]
    uniq, inv = np.unique(packed, return_inverse=True)
    ucol = np.stack([(uniq >> 16) & 255, (uniq >> 8) & 255, uniq & 255], -1).astype(np.float32) / 255
    out = np.empty_like(ucol)
    step = 262144
    for i in range(0, len(ucol), step):
        out[i:i + step] = transform_rgb(ucol[i:i + step], s)
    new = np.round(out[inv] * 255).astype(np.uint8).reshape(arr.shape[0], arr.shape[1], 3)
    if keep is not None and keep.any():
        new[keep] = arr[..., :3][keep]
    if has_alpha:
        new = np.dstack([new, arr[..., 3]])
    status = "converted" if keep is None or not keep.any() else "converted (kept %d photo panel(s))" % \
        (ndimage.label(keep)[1])
    return Image.fromarray(new, "RGBA" if has_alpha else "RGB"), status


def encode(img, fmt, info=None):
    """Save losslessly (PNG/TIFF) unless JPEG is asked for; keeps DPI and colour profile."""
    buf = io.BytesIO()
    extra = {}
    for k in ("dpi", "icc_profile"):
        if info and info.get(k):
            extra[k] = info[k]
    if "dpi" in extra:
        extra["dpi"] = tuple(float(v) for v in extra["dpi"])
    if fmt == "JPEG":
        img.convert("RGB").save(buf, "JPEG", quality=98, subsampling=0, **extra)
    elif fmt == "TIFF":
        img.save(buf, "TIFF", compression="tiff_lzw", **extra)
    else:
        img.save(buf, "PNG", compress_level=6, **extra)
    return buf.getvalue()


# ----------------------------------------------------------------------------
# SVG and metafiles
# ----------------------------------------------------------------------------

_COLOUR_RE = re.compile(
    r"(?P<pre>(?:fill|stroke|stop-color|flood-color|lighting-color|color)\s*(?:=\s*[\"']|:)\s*)"
    r"(?P<col>#[0-9a-fA-F]{6}\b|#[0-9a-fA-F]{3}\b|rgba?\([^)]*\)|[a-zA-Z]+)", re.I)
_DATAURI_RE = re.compile(r"data:image/(png|jpe?g|gif);base64,([A-Za-z0-9+/=\s]+)")
_SKIP_WORDS = {"none", "currentcolor", "inherit", "transparent", "url", "context"}


def transform_svg(text, s):
    def repl(m):
        col = m.group("col")
        if col.lower() in _SKIP_WORDS:
            return m.group(0)
        try:
            rgba = ImageColor.getrgb(col)
        except ValueError:
            return m.group(0)
        new = transform_hex("%02X%02X%02X" % rgba[:3], s)
        if col.lower().startswith("rgba") and len(rgba) == 4:
            return m.group("pre") + "rgba(%d,%d,%d,%s)" % (
                int(new[0:2], 16), int(new[2:4], 16), int(new[4:6], 16), col.split(",")[-1].strip(" )"))
        return m.group("pre") + "#" + new

    out = _COLOUR_RE.sub(repl, text)

    def repl_img(m):
        try:
            img = Image.open(io.BytesIO(base64.b64decode(m.group(2))))
            new, _ = transform_pil(img, s)
            if new is None:
                return m.group(0)
            return "data:image/png;base64," + base64.b64encode(encode(new, "PNG", img.info)).decode()
        except Exception:
            return m.group(0)

    out = _DATAURI_RE.sub(repl_img, out)
    # unstyled SVG shapes default to black fill -> make that default the inverted colour
    m = re.search(r"<svg\b[^>]*>", out)
    if m and not re.search(r"\sfill\s*=", m.group(0)) and "fill:" not in m.group(0):
        tag = m.group(0)
        out = out.replace(tag, tag[:-1].rstrip("/") + ' fill="#%s"' % transform_hex("000000", s)
                          + ("/>" if tag.endswith("/>") else ">"), 1)
    return out


def find_office():
    cands = [shutil.which("soffice"), shutil.which("libreoffice"),
             "/Applications/LibreOffice.app/Contents/MacOS/soffice",
             r"C:\Program Files\LibreOffice\program\soffice.exe"]
    for c in cands:
        if c and os.path.exists(c):
            return c
    return None


def find_inkscape():
    for c in [shutil.which("inkscape"), "/Applications/Inkscape.app/Contents/MacOS/inkscape",
              r"C:\Program Files\Inkscape\bin\inkscape.exe"]:
        if c and os.path.exists(c):
            return c
    return None


def rasterize_metafile(blob, ext, px_w, px_h):
    """EMF/WMF bytes -> PNG bytes, or None if no converter is available."""
    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "pic." + ext)
        with open(src, "wb") as f:
            f.write(blob)
        office = find_office()
        if office:
            filt = 'png:draw_png_Export:{"PixelWidth":{"type":"long","value":"%d"},' \
                   '"PixelHeight":{"type":"long","value":"%d"}}' % (px_w, px_h)
            prof = "file://" + os.path.join(d, "profile").replace("\\", "/")
            subprocess.run([office, "-env:UserInstallation=" + prof, "--headless",
                            "--convert-to", filt, "--outdir", d, src],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
            png = os.path.join(d, "pic.png")
            if os.path.exists(png):
                return open(png, "rb").read()
        ink = find_inkscape()
        if ink:
            png = os.path.join(d, "pic.png")
            subprocess.run([ink, src, "--export-type=png", "--export-filename=" + png,
                            "--export-width=%d" % px_w],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
            if os.path.exists(png):
                return open(png, "rb").read()
    return None


# ----------------------------------------------------------------------------
# EMF / WMF: recolour the vector records in place (no rasterising)
# ----------------------------------------------------------------------------

class _HexCache(dict):
    def __init__(self, s):
        super().__init__()
        self.s = s

    def __missing__(self, key):
        v = transform_hex(key, self.s)
        self[key] = v
        return v


def _swap_rgb(buf, o, cache):
    """COLORREF / RGB triple stored as R,G,B at buf[o:o+3]."""
    if o + 3 <= len(buf):
        buf[o:o + 3] = bytes.fromhex(cache[buf[o:o + 3].hex().upper()])


def _swap_bgr(buf, o, cache):
    """ARGB / RGBQUAD stored as B,G,R at buf[o:o+3]."""
    if o + 3 <= len(buf):
        bgr = buf[o:o + 3]
        new = bytes.fromhex(cache["%02X%02X%02X" % (bgr[2], bgr[1], bgr[0])])
        buf[o:o + 3] = bytes((new[2], new[1], new[0]))


def _dib_transform(buf, off_bmi, off_bits, cb_bits, s, cache):
    """Recolour a device-independent bitmap stored in buf (in place, same size)."""
    import struct
    if off_bmi <= 0 or off_bmi + 40 > len(buf):
        return
    bi_size, w, h, _, bpp, comp = struct.unpack_from("<IiiHHI", buf, off_bmi)
    clr_used = struct.unpack_from("<I", buf, off_bmi + 32)[0] if bi_size >= 36 else 0
    if bpp <= 8:
        n = clr_used or (1 << bpp)
        pal = off_bmi + bi_size + (12 if comp == 3 and bi_size == 40 else 0)
        for i in range(n):
            _swap_bgr(buf, pal + 4 * i, cache)
        return
    if bpp not in (24, 32) or comp not in (0, 3) or off_bits <= 0:
        return
    w, h = abs(w), abs(h)
    stride = ((w * bpp + 31) // 32) * 4
    if off_bits + stride * h > len(buf):
        return
    raw = np.frombuffer(bytes(buf[off_bits:off_bits + stride * h]), np.uint8).reshape(h, stride)
    px = raw[:, :w * bpp // 8].reshape(h, w, bpp // 8)
    img = Image.fromarray(np.ascontiguousarray(px[..., 2::-1]), "RGB")
    new, _ = transform_pil(img, s)
    if new is None:
        return
    out = raw.copy()
    view = out[:, :w * bpp // 8].reshape(h, w, bpp // 8)
    view[..., :3] = np.asarray(new)[..., ::-1]
    buf[off_bits:off_bits + stride * h] = out.tobytes()


def _u32(buf, o):
    import struct
    return struct.unpack_from("<I", buf, o)[0] if o + 4 <= len(buf) else 0


# EMF DIB records: record type -> offset of (offBmi, cbBmi, offBits, cbBits)
_EMF_DIB = {81: 48, 80: 48, 76: 84, 77: 84, 114: 84, 78: 84, 116: 84, 94: 16, 95: 12}
_STOCK_SWAP = {0: 4, 4: 0, 1: 3, 3: 1, 6: 7, 7: 6}


def _emfplus_brush(buf, o, end, cache):
    """Recolour an EmfPlusBrush starting at o. Returns nothing."""
    btype = _u32(buf, o + 4)
    d = o + 8
    if btype == 0:
        _swap_bgr(buf, d, cache)
    elif btype == 1:
        _swap_bgr(buf, d + 4, cache)
        _swap_bgr(buf, d + 8, cache)
    elif btype == 3:
        _swap_bgr(buf, d + 8, cache)
        n = _u32(buf, d + 20)
        for i in range(min(n, 4096)):
            _swap_bgr(buf, d + 24 + 4 * i, cache)
    elif btype == 4:
        flags = _u32(buf, d)
        _swap_bgr(buf, d + 24, cache)
        _swap_bgr(buf, d + 28, cache)
        q = d + 40
        if flags & 0x02:
            q += 24
        if flags & 0x04:
            n = _u32(buf, q)
            cols = q + 4 + 4 * n
            for i in range(min(n, 4096)):
                if cols + 4 * i + 4 <= end:
                    _swap_bgr(buf, cols + 4 * i, cache)


def _emfplus_pen(buf, o, end, cache):
    d = o + 8                      # skip Version, Type
    flags = _u32(buf, d)
    q = d + 12                     # PenDataFlags, PenUnit, PenWidth
    for bit, size in ((0x01, 24), (0x02, 4), (0x04, 4), (0x08, 4), (0x10, 4), (0x20, 4), (0x40, 4), (0x80, 4)):
        if flags & bit:
            q += size
    if flags & 0x100:
        q += 4 + 4 * _u32(buf, q)
    if flags & 0x200:
        q += 4
    if flags & 0x400:
        q += 4 + 4 * _u32(buf, q)
    if flags & 0x800:
        q += 4 + _u32(buf, q)
    if flags & 0x1000:
        q += 4 + _u32(buf, q)
    if q < end:
        _emfplus_brush(buf, q, end, cache)


def _emfplus_image(rec, s):
    """EmfPlusObject image record (bytearray, whole record). Returns replacement record or None."""
    import struct
    d = 12
    if _u32(rec, d + 4) != 1:      # not a bitmap
        return None
    w, h, stride, pfmt, btype = struct.unpack_from("<iiiII", rec, d + 8)
    data = d + 28
    if btype == 1:                 # compressed (PNG / JPEG / ...)
        try:
            img = Image.open(io.BytesIO(bytes(rec[data:])))
            img.load()
        except Exception:
            return None
        new, _ = transform_pil(img, s)
        if new is None:
            return None
        payload = encode(new, "PNG")
        payload += b"\0" * (-len(payload) % 4)
        out = bytearray(rec[:data]) + payload
        struct.pack_into("<II", out, 4, len(out), len(out) - 12)
        return out
    if pfmt in (0x0026200A, 0x00022009, 0x000E200B, 0x00021808) and stride > 0:
        bpp = 3 if pfmt == 0x00021808 else 4
        if data + stride * h > len(rec):
            return None
        raw = np.frombuffer(bytes(rec[data:data + stride * h]), np.uint8).reshape(h, stride)
        px = raw[:, :w * bpp].reshape(h, w, bpp)
        new, _ = transform_pil(Image.fromarray(np.ascontiguousarray(px[..., 2::-1]), "RGB"), s)
        if new is None:
            return None
        out_raw = raw.copy()
        out_raw[:, :w * bpp].reshape(h, w, bpp)[..., :3] = np.asarray(new)[..., ::-1]
        rec[data:data + stride * h] = out_raw.tobytes()
        return rec
    return None


_PLUS_FILL_S = {0x400A, 0x400C, 0x400E, 0x4010, 0x4013, 0x4014, 0x4016, 0x401C, 0x4036}


def _emfplus_block(block, s, cache):
    """block: bytes of EMF+ records. Returns new bytes."""
    import struct
    out = bytearray()
    o = 0
    while o + 12 <= len(block):
        rtype, flags, size, dsize = struct.unpack_from("<HHII", block, o)
        if size < 12 or o + size > len(block):
            out += block[o:]
            break
        rec = bytearray(block[o:o + size])
        try:
            if rtype == 0x4008 and not flags & 0x8000:
                otype = (flags >> 8) & 0x7F
                if otype == 1:
                    _emfplus_brush(rec, 12, size, cache)
                elif otype == 2:
                    _emfplus_pen(rec, 12, size, cache)
                elif otype == 5:
                    rec = _emfplus_image(rec, s) or rec
            elif rtype == 0x4009:
                _swap_bgr(rec, 12, cache)
            elif rtype in _PLUS_FILL_S and flags & 0x8000:
                _swap_bgr(rec, 12, cache)
        except Exception:
            pass
        out += rec
        o += size
    return bytes(out)


def _emf_record(rtype, payload):
    import struct
    return struct.pack("<II", rtype, 8 + len(payload)) + payload


def recolour_emf(blob, s):
    import struct
    cache = _HexCache(s)
    if len(blob) < 88 or _u32(blob, 0) != 1 or blob[40:44] != b" EMF":
        raise ValueError("not an EMF file")
    records = []
    o = 0
    while o + 8 <= len(blob):
        rtype, size = struct.unpack_from("<II", blob, o)
        if size < 8 or o + size > len(blob):
            break
        records.append(bytearray(blob[o:o + size]))
        o += size
        if rtype == 14:  # EOF
            break

    for idx, rec in enumerate(records):
        rtype = _u32(rec, 0)
        try:
            if rtype in (24, 25):
                _swap_rgb(rec, 8, cache)
            elif rtype == 38:
                _swap_rgb(rec, 24, cache)
            elif rtype == 39:
                _swap_rgb(rec, 16, cache)
            elif rtype == 15:
                _swap_rgb(rec, 16, cache)
            elif rtype == 37:
                ih = _u32(rec, 8)
                if ih & 0x80000000 and (ih & 0x7FFFFFFF) in _STOCK_SWAP:
                    struct.pack_into("<I", rec, 8, 0x80000000 | _STOCK_SWAP[ih & 0x7FFFFFFF])
            elif rtype == 118:
                nv = _u32(rec, 24)
                for i in range(min(nv, 65536)):
                    v = 36 + 16 * i + 8
                    if v + 6 <= len(rec):
                        r_, g_, b_ = (rec[v + 1], rec[v + 3], rec[v + 5])
                        new = bytes.fromhex(cache["%02X%02X%02X" % (r_, g_, b_)])
                        for k in range(3):
                            struct.pack_into("<H", rec, v + 2 * k, new[k] << 8)
            if rtype == 95:
                _swap_rgb(rec, 40, cache)
            if rtype in (76, 77, 114, 78, 116):
                _swap_rgb(rec, 76, cache)          # crBkColorSrc
                if rtype == 116:
                    _swap_rgb(rec, 40, cache)      # transparent colour key
            if rtype in _EMF_DIB:
                p = _EMF_DIB[rtype]
                off_bmi, cb_bmi, off_bits, cb_bits = struct.unpack_from("<IIII", rec, p)
                if cb_bmi:
                    _dib_transform(rec, off_bmi, off_bits, cb_bits, s, cache)
            if rtype == 70 and rec[12:16] == b"EMF+":
                dsize = _u32(rec, 8)
                block = bytes(rec[16:12 + dsize])
                new_block = _emfplus_block(block, s, cache)
                if new_block != block:
                    payload = struct.pack("<I", 4 + len(new_block)) + b"EMF+" + new_block
                    payload += b"\0" * (-len(payload) % 4)
                    records[idx] = bytearray(_emf_record(70, payload))
        except Exception:
            pass

    # GDI defaults are black text/pens and white brushes: set inverted defaults up front
    defaults = [
        _emf_record(24, bytes.fromhex(cache["000000"]) + b"\0"),
        _emf_record(25, bytes.fromhex(cache["FFFFFF"]) + b"\0"),
        _emf_record(37, struct.pack("<I", 0x80000000 | 6)),
        _emf_record(37, struct.pack("<I", 0x80000000 | 4)),
    ]
    at = 1
    if len(records) > 1 and _u32(records[1], 0) == 70 and records[1][12:16] == b"EMF+":
        at = 2   # the EMF+ header comment must stay first
    records[at:at] = [bytearray(r) for r in defaults]
    out = bytearray(b"".join(records))
    struct.pack_into("<II", out, 48, len(out), len(records))
    return bytes(out)


def recolour_wmf(blob, s):
    import struct
    cache = _HexCache(s)
    buf = bytearray(blob)
    start = 22 if _u32(buf, 0) == 0x9AC6CDD7 else 0
    hdr_words = struct.unpack_from("<H", buf, start + 2)[0]
    o = start + hdr_words * 2
    records = []
    while o + 6 <= len(buf):
        size = _u32(buf, o) * 2
        if size < 6 or o + size > len(buf):
            break
        rec = bytearray(buf[o:o + size])
        fn = struct.unpack_from("<H", rec, 4)[0]
        try:
            if fn in (0x0201, 0x0209, 0x041F):
                _swap_rgb(rec, 6, cache)
            elif fn == 0x02FA:
                _swap_rgb(rec, 12, cache)
            elif fn == 0x02FC:
                _swap_rgb(rec, 8, cache)
            dib = {0x0F43: 28, 0x0B41: 26, 0x0940: 22, 0x0D33: 24, 0x0142: 10}.get(fn)
            if dib and size > dib + 40:
                bi_size, _, _, _, bpp, comp = struct.unpack_from("<IiiHHI", rec, dib)
                clr_used = _u32(rec, dib + 32)
                pal = (clr_used or (1 << bpp)) * 4 if bpp <= 8 else 0
                bits = dib + bi_size + pal + (12 if comp == 3 and bi_size == 40 else 0)
                _dib_transform(rec, dib, bits, len(rec) - bits, s, cache)
        except Exception:
            pass
        records.append(rec)
        o += size
        if fn == 0:
            break
    pre = [struct.pack("<IH", 5, 0x0209) + bytes.fromhex(cache["000000"]) + b"\0",
           struct.pack("<IH", 5, 0x0201) + bytes.fromhex(cache["FFFFFF"]) + b"\0"]
    body = b"".join(pre) + b"".join(records)
    header = bytearray(buf[start:start + hdr_words * 2])
    total_words = (len(header) + len(body)) // 2
    struct.pack_into("<I", header, 6, total_words)
    struct.pack_into("<I", header, 12, max(_u32(header, 12), 5))
    return bytes(buf[:start]) + bytes(header) + body


# ----------------------------------------------------------------------------
# DrawingML colours
# ----------------------------------------------------------------------------

COLOUR_TAGS = {"srgbClr", "schemeClr", "sysClr", "prstClr", "scrgbClr", "hslClr"}
TEXT_TAGS = {"rPr", "defRPr", "endParaRPr", "fontRef", "buClr"}
SKIP_TAGS = {"effectLst", "effectRef", "effectDag", "bg", "clrScheme", "extLst"}
FILL_TAGS = {"solidFill", "gradFill", "noFill", "blipFill", "pattFill", "grpFill"}
KEEP_MODS = {"alpha", "alphaMod", "alphaOff"}


def _hsl_mod(rgb, mods):
    r, g, b = rgb
    h, l, s_ = colorsys.rgb_to_hls(r, g, b)
    for tag, v in mods:
        v = v / 100000
        if tag == "lumMod":
            l *= v
        elif tag == "lumOff":
            l += v
        elif tag == "satMod":
            s_ *= v
        elif tag == "satOff":
            s_ += v
        elif tag == "hueOff":
            h = (h + v * 100000 / 60000 / 360) % 1
    r, g, b = colorsys.hls_to_rgb(h, min(max(l, 0), 1), min(max(s_, 0), 1))
    return [r, g, b]


class Theme:
    def __init__(self, theme_el, clr_map):
        self.colours = {}
        if theme_el is not None:
            scheme = theme_el.find(".//" + a("clrScheme"))
            for child in scheme if scheme is not None else []:
                c = child[0] if len(child) else None
                if c is None:
                    continue
                if local(c) == "srgbClr":
                    self.colours[local(child)] = c.get("val").upper()
                elif local(c) == "sysClr":
                    self.colours[local(child)] = (c.get("lastClr") or
                                                  ("FFFFFF" if c.get("val") == "window" else "000000")).upper()
        self.colours.setdefault("dk1", "000000")
        self.colours.setdefault("lt1", "FFFFFF")
        self.clr_map = clr_map

    def scheme(self, name):
        name = self.clr_map.get(name, name)
        return self.colours.get(name)


def resolve_colour(el, theme):
    """DrawingML colour element -> 'RRGGBB' (modifiers applied) or None."""
    tag = local(el)
    if tag == "srgbClr":
        base = el.get("val")
    elif tag == "schemeClr":
        base = theme.scheme(el.get("val"))
    elif tag == "sysClr":
        base = el.get("lastClr") or ("FFFFFF" if el.get("val") == "window" else "000000")
    elif tag == "prstClr":
        try:
            base = "%02X%02X%02X" % ImageColor.getrgb(el.get("val"))[:3]
        except ValueError:
            base = None
    elif tag == "scrgbClr":
        lin = [int(el.get(k, 0)) / 100000 for k in "rgb"]
        base = "%02X%02X%02X" % tuple(int(round(float(_to_srgb(np.array(v))) * 255)) for v in lin)
    elif tag == "hslClr":
        r, g, b = colorsys.hls_to_rgb(int(el.get("hue", 0)) / 21600000, int(el.get("lum", 0)) / 100000,
                                      int(el.get("sat", 0)) / 100000)
        base = "%02X%02X%02X" % tuple(int(round(v * 255)) for v in (r, g, b))
    else:
        base = None
    if not base:
        return None
    rgb = [int(base[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    hsl_mods = []
    for m in el:
        mt, v = local(m), m.get("val")
        if v is None:
            continue
        v = int(v)
        if mt in ("lumMod", "lumOff", "satMod", "satOff", "hueOff"):
            hsl_mods.append((mt, v))
        elif mt == "tint":
            rgb = _hsl_mod(rgb, hsl_mods); hsl_mods = []
            rgb = [c * v / 100000 + (1 - v / 100000) for c in rgb]
        elif mt == "shade":
            rgb = _hsl_mod(rgb, hsl_mods); hsl_mods = []
            rgb = [c * v / 100000 for c in rgb]
    rgb = _hsl_mod(rgb, hsl_mods) if hsl_mods else rgb
    return "%02X%02X%02X" % tuple(int(round(min(max(c, 0), 1) * 255)) for c in rgb)


def set_colour(el, hexval):
    new = etree.Element(a("srgbClr"), val=hexval)
    for m in el:
        if local(m) in KEEP_MODS:
            new.append(m)
    el.getparent().replace(el, new)


def _fill_colour(fill_el, theme):
    tag = local(fill_el)
    if tag == "noFill":
        return None
    if tag == "solidFill" and len(fill_el):
        return resolve_colour(fill_el[0], theme)
    if tag in ("gradFill", "pattFill"):
        cols = [resolve_colour(c, theme) for c in fill_el.iter() if local(c) in COLOUR_TAGS]
        cols = [c for c in cols if c]
        if cols:
            arr = np.array([[int(c[i:i + 2], 16) for i in (0, 2, 4)] for c in cols]).mean(0)
            return "%02X%02X%02X" % tuple(int(v) for v in arr)
    return "unknown"


def container_fill(container, theme):
    """Original fill colour behind text in this shape / table cell: hex, None (transparent) or 'unknown'."""
    if container is None:
        return None
    props = None
    for child in container:
        if local(child) in ("spPr", "tcPr"):
            props = child
    if props is not None:
        for child in props:
            if local(child) in FILL_TAGS:
                return _fill_colour(child, theme)
    for child in container:
        if local(child) == "style":
            for ref in child:
                if local(ref) == "fillRef" and ref.get("idx", "0") != "0" and len(ref):
                    return resolve_colour(ref[0], theme)
    return None


def recolour_tree(root, theme, s, counts):
    """Re-colour all explicit colours in an XML tree for a black background."""
    bg_hex = "%02X%02X%02X" % s.bg
    fill_cache = {}

    def new_bg_for(container):
        key = id(container)
        if key not in fill_cache:
            f = container_fill(container, theme)
            if f is None:
                fill_cache[key] = bg_hex
            elif f == "unknown":
                fill_cache[key] = "unknown"
            else:
                fill_cache[key] = transform_hex(f, s)
        return fill_cache[key]

    def container_of(el):
        for anc in el.iterancestors():
            if local(anc) in ("sp", "tc"):
                return anc
        return None

    changes = []
    for el in root.iter():
        if local(el) not in COLOUR_TAGS:
            continue
        if el.get("val") == "phClr":
            continue
        ancestors = {local(x) for x in el.iterancestors()}
        if ancestors & SKIP_TAGS:
            continue
        orig = resolve_colour(el, theme)
        if orig is None:
            continue
        new = transform_hex(orig, s)
        if ancestors & TEXT_TAGS and "highlight" not in ancestors:
            behind = new_bg_for(container_of(el))
            if behind == "unknown" or contrast(orig, behind) > contrast(new, behind):
                new = orig
            if behind != "unknown" and rel_luminance(behind) < 0.05:
                new = lift_lightness(new, MIN_TEXT_L)
        changes.append((el, new))

    # inherited (theme) text in filled shapes: pin it if the fill stays light/colourful
    default_text = theme.scheme("tx1") or "000000"
    for container in root.iter():
        if local(container) != "sp":
            continue
        f = container_fill(container, theme)
        if f in (None, "unknown"):
            continue
        if any(local(c) == "style" and any(local(r) == "fontRef" and len(r) for r in c) for c in container):
            continue
        behind = transform_hex(f, s)
        if contrast(default_text, behind) <= contrast(transform_hex(default_text, s), behind):
            continue
        for run in container.iter(a("r"), a("fld")):
            rpr = run.find(a("rPr"))
            if rpr is None:
                rpr = etree.Element(a("rPr"))
                run.insert(0, rpr)
            if any(local(c) in FILL_TAGS for c in rpr):
                continue
            fill = etree.Element(a("solidFill"))
            etree.SubElement(fill, a("srgbClr"), val=default_text)
            ln = rpr.find(a("ln"))
            rpr.insert(0 if ln is None else list(rpr).index(ln) + 1, fill)

    for el, new in changes:
        set_colour(el, new)
    counts["colours"] += len(changes)


# ----------------------------------------------------------------------------
# Presentation handling
# ----------------------------------------------------------------------------

CLR_MAP_ATTRS = ["bg1", "tx1", "bg2", "tx2", "accent1", "accent2", "accent3", "accent4",
                 "accent5", "accent6", "hlink", "folHlink"]


def master_clr_map(master):
    cm = master._element.find("{%s}clrMap" % P_NS)
    d = dict(cm.attrib) if cm is not None else {}
    defaults = dict(bg1="lt1", tx1="dk1", bg2="lt2", tx2="dk2", hlink="hlink", folHlink="folHlink",
                    **{"accent%d" % i: "accent%d" % i for i in range(1, 7)})
    defaults.update(d)
    return defaults


def slide_clr_map(slide):
    cm = master_clr_map(slide.slide_layout.slide_master)
    ovr = slide._element.find("{%s}clrMapOvr" % P_NS)
    if ovr is not None:
        o = ovr.find(a("overrideClrMapping"))
        if o is not None:
            cm.update(o.attrib)
    return cm


def set_dark_clr_map(slide, cm):
    """Swap bg/tx so inherited theme text becomes light on this slide only."""
    dark = dict(cm)
    dark["bg1"], dark["tx1"] = cm["tx1"], cm["bg1"]
    dark["bg2"], dark["tx2"] = cm["tx2"], cm["bg2"]
    sld = slide._element
    ovr = sld.find("{%s}clrMapOvr" % P_NS)
    if ovr is None:
        ovr = etree.Element("{%s}clrMapOvr" % P_NS)
        csld = sld.find("{%s}cSld" % P_NS)
        csld.addnext(ovr)
    for c in list(ovr):
        ovr.remove(c)
    etree.SubElement(ovr, a("overrideClrMapping"), **{k: dark[k] for k in CLR_MAP_ATTRS})


def set_background(slide, rgb):
    csld = slide._element.find("{%s}cSld" % P_NS)
    old = csld.find("{%s}bg" % P_NS)
    if old is not None:
        csld.remove(old)
    bg = etree.Element("{%s}bg" % P_NS)
    bgpr = etree.SubElement(bg, "{%s}bgPr" % P_NS)
    fill = etree.SubElement(bgpr, a("solidFill"))
    etree.SubElement(fill, a("srgbClr"), val="%02X%02X%02X" % rgb)
    etree.SubElement(bgpr, a("effectLst"))
    csld.insert(0, bg)


def _alt_text(el):
    for anc in el.iterancestors():
        if local(anc) in ("pic", "sp"):
            for c in anc.iter():
                if local(c) == "cNvPr":
                    return ((c.get("descr") or "") + " " + (c.get("title") or "")).lower()
            break
    return ""


def _target_px(el, native_size, dpi):
    """Pixel size to rasterise a metafile at, from the picture's size on the slide."""
    for anc in el.iterancestors():
        if local(anc) in ("pic", "sp"):
            ext = anc.find(".//" + a("ext"))
            if ext is not None and ext.get("cx"):
                w = int(ext.get("cx")) / 914400 * dpi
                h = int(ext.get("cy")) / 914400 * dpi
                return max(64, int(w)), max(64, int(h))
            break
    return native_size


class Converter:
    def __init__(self, prs, s, log):
        self.prs = prs
        self.s = s
        self.log = log
        self.cache = {}  # (original partname, force) -> (new blob, ext) or None
        self.warnings = set()

    def _new_part(self, host_part, blob, ext):
        from pptx.opc.package import Part
        from pptx.opc.packuri import PackURI
        types = {"svg": "image/svg+xml", "emf": "image/x-emf", "wmf": "image/x-wmf"}
        if ext in types:
            pn = self.prs.part.package.next_partname("/ppt/media/image%d." + ext)
            part = Part(PackURI(pn), types[ext], self.prs.part.package, blob)
            return host_part.relate_to(part, RT_IMAGE)
        _, rid = host_part.get_or_add_image_part(io.BytesIO(blob))
        return rid

    def convert_image_part(self, img_part, el, force):
        key = (str(img_part.partname), force)
        if key in self.cache:
            return self.cache[key]
        blob, ct = img_part.blob, img_part.content_type
        ext = str(img_part.partname).rsplit(".", 1)[-1].lower()
        result, status = None, "skipped"
        try:
            if ct == "image/svg+xml" or ext == "svg":
                result = (transform_svg(blob.decode("utf-8"), self.s).encode("utf-8"), "svg")
                status = "svg recoloured"
            elif ext in ("emf", "wmf", "emz", "wmz") or "metafile" in ct or "wmf" in ct or "emf" in ct:
                if ext in ("emz", "wmz"):
                    import gzip
                    blob = gzip.decompress(blob)
                kind = "emf" if blob[40:44] == b" EMF" else "wmf"
                if self.s.rasterize_metafiles:
                    w, h = _target_px(el, (1600, 1200), self.s.dpi)
                    png = rasterize_metafile(blob, kind, w, h)
                    if png is None:
                        self.warnings.add("--rasterize-metafiles needs LibreOffice or Inkscape; "
                                          "EMF/WMF pictures were left unchanged.")
                        status = "metafile (no converter)"
                    else:
                        new, status = transform_pil(Image.open(io.BytesIO(png)), self.s, force)
                        result = (encode(new, "PNG"), "png") if new is not None else None
                        status = "%s rasterised at %d dpi, %s" % (kind.upper(), self.s.dpi, status)
                else:
                    fn = recolour_emf if kind == "emf" else recolour_wmf
                    result = (fn(blob, self.s), kind)
                    status = "%s recoloured (still vector)" % kind.upper()
            else:
                img = Image.open(io.BytesIO(blob))
                fmt = img.format
                if getattr(img, "n_frames", 1) > 1:
                    status = "animated (skipped)"
                else:
                    new, status = transform_pil(img, self.s, force)
                    if new is not None:
                        out_fmt = "JPEG" if (fmt == "JPEG" and self.s.keep_jpeg) else "PNG"
                        result = (encode(new, out_fmt, img.info), "jpg" if out_fmt == "JPEG" else "png")
                        status += " %dx%d px" % new.size
        except Exception as e:  # never let one odd picture stop the deck
            status = "error: %s" % e
        self.cache[key] = (result, status)
        return result, status

    def process_blips(self, host_part, root, counts):
        rels_to_check = set()
        for blip in list(root.iter()):
            if local(blip) not in ("blip", "svgBlip"):
                continue
            rid = blip.get(R_EMBED)
            if not rid or rid not in host_part.rels:
                continue
            alt = _alt_text(blip)
            if "#keep" in alt:
                counts["images_kept"] += 1
                continue
            img_part = host_part.rels[rid].target_part
            result, status = self.convert_image_part(img_part, blip, "#force" in alt)
            self.log("    picture %-28s %s" % (str(img_part.partname).rsplit("/", 1)[-1], status))
            if result is None:
                counts["images_kept"] += 1
                continue
            new_rid = self._new_part(host_part, *result)
            blip.set(R_EMBED, new_rid)
            rels_to_check.add(rid)
            counts["images"] += 1
        for rid in rels_to_check:
            host_part.drop_rel(rid)

    def process_slide(self, slide, counts):
        cm = slide_clr_map(slide)
        master = slide.slide_layout.slide_master
        theme_el = etree.fromstring(master.part.part_related_by(RT_THEME).blob)
        theme = Theme(theme_el, cm)

        recolour_tree(slide._element.find("{%s}cSld" % P_NS), theme, self.s, counts)
        self.process_blips(slide.part, slide._element, counts)

        for rel in list(slide.part.rels.values()):
            if rel.is_external:
                continue
            if rel.reltype == RT_CHART:
                part = rel.target_part
                recolour_tree(part._element, theme, self.s, counts)
                self.process_blips(part, part._element, counts)
                counts["charts"] += 1
            elif rel.reltype == RT_DGM_DRAWING:
                part = rel.target_part
                root = etree.fromstring(part.blob)
                recolour_tree(root, theme, self.s, counts)
                part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                counts["smartart"] += 1

        set_dark_clr_map(slide, cm)
        set_background(slide, self.s.bg)

    def process_layout_shapes(self, part_obj, counts, theme):
        """Non-placeholder decoration on layouts/masters (logos, bars, lines)."""
        tree = part_obj._element.find("{%s}cSld/{%s}spTree" % (P_NS, P_NS))
        for shape in list(tree):
            if shape.find(".//{%s}ph" % P_NS) is not None:
                continue
            recolour_tree(shape, theme, self.s, counts)
        self.process_blips(part_obj.part, tree, counts)


def parse_slides(spec, n):
    if not spec:
        return list(range(1, n + 1))
    out = []
    for chunk in spec.split(","):
        chunk = chunk.strip()
        if "-" in chunk:
            lo, hi = chunk.split("-")
            out.extend(range(int(lo), int(hi or n) + 1))
        elif chunk:
            out.append(int(chunk))
    return [i for i in out if 1 <= i <= n]


def convert_pptx(src, dst, s, slides=None, layouts=False, quiet=False):
    from pptx import Presentation
    log = (lambda *x: None) if quiet else print
    prs = Presentation(src)
    conv = Converter(prs, s, log)
    picked = parse_slides(slides, len(prs.slides))
    totals = dict(colours=0, images=0, images_kept=0, charts=0, smartart=0)
    for i, slide in enumerate(prs.slides, start=1):
        if i not in picked:
            continue
        counts = dict.fromkeys(totals, 0)
        log("slide %d" % i)
        conv.process_slide(slide, counts)
        log("    %d colours, %d pictures converted, %d left as-is" %
            (counts["colours"], counts["images"], counts["images_kept"]))
        for k in totals:
            totals[k] += counts[k]
    if layouts:
        done = set()
        for i, slide in enumerate(prs.slides, start=1):
            if i not in picked:
                continue
            for obj in (slide.slide_layout, slide.slide_layout.slide_master):
                if id(obj.part) in done:
                    continue
                done.add(id(obj.part))
                master = slide.slide_layout.slide_master
                theme = Theme(etree.fromstring(master.part.part_related_by(RT_THEME).blob),
                              master_clr_map(master))
                log("layout/master %s" % obj.part.partname)
                conv.process_layout_shapes(obj, totals, theme)
    prs.save(dst)
    for w in sorted(conv.warnings):
        print("WARNING:", w)
    log("saved %s  (%d slides, %d colours, %d pictures converted)" %
        (dst, len(picked), totals["colours"], totals["images"]))


def convert_image_file(src, dst, s, force=False):
    ext = src.rsplit(".", 1)[-1].lower()
    if ext == "svg":
        with open(src, encoding="utf-8") as f:
            out = transform_svg(f.read(), s)
        with open(dst, "w", encoding="utf-8") as f:
            f.write(out)
        print("saved", dst)
        return
    if ext in ("emf", "wmf"):
        blob = open(src, "rb").read()
        with open(dst, "wb") as f:
            f.write(recolour_emf(blob, s) if blob[40:44] == b" EMF" else recolour_wmf(blob, s))
        print("saved", dst, "(vector)")
        return
    else:
        img = Image.open(src)
        fmt = {"TIFF": "TIFF"}.get(img.format, "PNG")
        if img.format == "JPEG":
            fmt = "JPEG" if s.keep_jpeg else "PNG"
            if fmt == "PNG":
                dst = dst.rsplit(".", 1)[0] + ".png"
    new, status = transform_pil(img, s, force)
    if new is None:
        print("%s looks like it already has a dark background; use --force to convert anyway" % src)
        return
    with open(dst, "wb") as f:
        f.write(encode(new, fmt, img.info))
    print("saved %s (%s)" % (dst, status))


def main(argv=None):
    p = argparse.ArgumentParser(description="Convert PowerPoint slides / figure images for a black background.",
                                formatter_class=argparse.RawDescriptionHelpFormatter, epilog=__doc__)
    p.add_argument("inputs", nargs="+", help=".pptx file(s) or image files (png, jpg, tif, svg, emf)")
    p.add_argument("-o", "--output", help="output file (single input only); default adds _black")
    p.add_argument("--slides", help='slides to convert, e.g. "2,4-7" (default: all)')
    p.add_argument("--method", choices=["oklab", "hsl", "invert"], default="oklab",
                   help="oklab (default): neutrals invert, colours keep hue and stay bright; "
                        "hsl: invert HSL lightness (like ezreverse); invert: plain negative")
    p.add_argument("--brighten", type=float, default=0.0, help="0-1, lift mid-tones (default 0)")
    p.add_argument("--saturation", type=float, default=1.0, help="colour boost, e.g. 1.2 (default 1)")
    p.add_argument("--protect", type=float, default=1.0,
                   help="0-1, how much vivid colours keep their brightness instead of inverting (default 1)")
    p.add_argument("--black-point", type=float, default=0.0,
                   help="0-0.2: near-black output snaps to pure black, e.g. 0.05 cleans JPEG noise (default 0)")
    p.add_argument("--bg", default="000000", help="background colour hex (default 000000)")
    p.add_argument("--keep-jpeg", action="store_true",
                   help="save converted JPEGs as JPEG (smaller files). Default: lossless PNG")
    p.add_argument("--rasterize-metafiles", action="store_true",
                   help="turn EMF/WMF into PNG at --dpi instead of recolouring them as vectors")
    p.add_argument("--dpi", type=int, default=600, help="resolution for --rasterize-metafiles (default 600)")
    p.add_argument("--no-photo-protect", action="store_true",
                   help="also invert dark photos / micrographs (normally left alone)")
    p.add_argument("--layouts", action="store_true",
                   help="also convert logos/lines on the slide layouts & master (affects every slide "
                        "using them; best when converting the whole deck)")
    p.add_argument("--force", action="store_true", help="image files: convert even if already dark")
    p.add_argument("-q", "--quiet", action="store_true")
    args = p.parse_args(argv)

    bg = args.bg.lstrip("#")
    s = Settings(method=args.method, protect=args.protect, brighten=args.brighten,
                 saturation=args.saturation, black_point=args.black_point,
                 bg=tuple(int(bg[i:i + 2], 16) for i in (0, 2, 4)), dpi=args.dpi,
                 protect_photos=not args.no_photo_protect, keep_jpeg=args.keep_jpeg,
                 rasterize_metafiles=args.rasterize_metafiles)
    if args.output and len(args.inputs) > 1:
        p.error("--output works with a single input")
    for src in args.inputs:
        root, ext = os.path.splitext(src)
        dst = args.output or root + "_black" + ext
        if ext.lower() in (".pptx", ".pptm", ".potx"):
            convert_pptx(src, dst, s, slides=args.slides, layouts=args.layouts, quiet=args.quiet)
        else:
            convert_image_file(src, dst, s, force=args.force)


if __name__ == "__main__":
    main()
